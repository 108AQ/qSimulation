"""Build qite-slides.pptx (python-pptx). Run from anywhere; writes next to the .tex."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
NAVY, BLUE, RED, GRAY, LIGHT = RGBColor(0x0B,0x1F,0x3A), RGBColor(0x00,0x33,0x99), RGBColor(0xA0,0,0), RGBColor(0x55,0x55,0x55), RGBColor(0xEE,0xF2,0xF8)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

def tb(s, x, y, w, h, paras, size=18, color=NAVY, bold=False, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = box.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, p in enumerate(paras if isinstance(paras, list) else [paras]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        t, kw = (p, {}) if isinstance(p, str) else p
        run = para.add_run(); run.text = t; para.alignment = align
        f = run.font; f.size = Pt(kw.get("size", size)); f.bold = kw.get("bold", bold)
        f.color.rgb = kw.get("color", color); f.name = kw.get("font", font)
        para.space_after = Pt(kw.get("after", 6))
    return box

def slide(title, n):
    s = prs.slides.add_slide(blank)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.05)); bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tb(s, 0.5, 0.18, 12.3, 0.8, title, size=30, color=RGBColor(255,255,255), bold=True, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 12.2, 7.05, 1, 0.3, str(n), size=11, color=GRAY, align=PP_ALIGN.RIGHT)
    return s

def bullets(s, items, x=0.6, y=1.4, w=12.1, h=5.5, size=20):
    paras = []
    for it in items:
        if isinstance(it, tuple): paras.append(("      – " + it[0], dict(size=size-3, color=GRAY, after=4)))
        else: paras.append(("•  " + it, dict(size=size, after=9)))
    return tb(s, x, y, w, h, paras, size=size)

def eq(s, text, x, y, w, h=0.7, size=24):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = LIGHT; r.line.fill.background()
    tb(s, x, y, w, h, text, size=size, color=BLUE, align=PP_ALIGN.CENTER, font="Cambria Math", anchor=MSO_ANCHOR.MIDDLE)

def table(s, rows, x, y, w, colw, size=14, rowh=0.5):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh*len(rows))).table
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = v
            for p in c.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size); run.font.name = "Calibri"; run.font.bold = (i == 0)
                    run.font.color.rgb = RGBColor(255,255,255) if i == 0 else NAVY
            c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else RGBColor(255,255,255))
    return t

n = 1
# 1 title
s = prs.slides.add_slide(blank); bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
tb(s, 0.8, 2.2, 11.7, 1.6, "Quantum Imaginary-Time Evolution (QITE)", size=44, color=RGBColor(255,255,255), bold=True)
tb(s, 0.8, 3.7, 11.7, 1.2, "Ground-state energies on a quantum computer — and how to benchmark it against Quantum Krylov on Quantinuum hardware", size=22, color=RGBColor(0xC8,0xD6,0xF0))
tb(s, 0.8, 6.2, 11.7, 0.5, "Prajjwal · 2026-10-08 · companion to quantum-imaginary-time-evolution.pdf", size=14, color=RGBColor(0xC8,0xD6,0xF0))

n += 1; s = slide("Goal and roadmap", n)
bullets(s, ["Goal: find the molecular ground-state energy  E₀ = min ⟨ψ|H|ψ⟩",
            "Idea: evolve in imaginary time, e^{−τH}, which filters out excited states",
            "Obstacle: e^{−τH} is not unitary, so it is not a circuit",
            "QITE: replace each small step by a unitary e^{−iΔτA}, with A fitted from measured Pauli expectation values",
            "Then: cost, pitfalls, comparison with Quantum Krylov, a worked H₂ check, and a benchmark plan"], size=22)

n += 1; s = slide("Why imaginary time finds the ground state", n)
eq(s, "|ψ(τ)⟩ ∝ e^{−τH}|ψ₀⟩ = e^{−τE₀}[ c₀|E₀⟩ + Σₙ>₀ cₙ e^{−τ(Eₙ−E₀)} |Eₙ⟩ ]", 0.5, 1.3, 7.3, 1.0, size=20)
bullets(s, ["Excited weight decays like e^{−2τ(Eₙ−E₀)}",
            "Needs overlap c₀ ≠ 0 (HF: |c₀|² ≈ 0.99 for H₂)",
            "E(τ) − E₀ = O(e^{−2τΔ}): the gap Δ sets the cost",
            "Only the gap within the starting symmetry sector matters",
            "Continuous power iteration"], x=0.5, y=2.6, w=7.3, size=19)
s.shapes.add_picture(str(ROOT/"figures"/"filter.png"), Inches(8.1), Inches(1.6), width=Inches(4.9))

n += 1; s = slide("The energy only goes down (exact ITE)", n)
eq(s, "d|ψ⟩/dτ = −(H − E)|ψ⟩          dE/dτ = −2 Var_ψ(H) ≤ 0", 0.8, 1.4, 11.7, 0.9, size=26)
bullets(s, ["Proof: dE/dτ = 2 Re⟨ψ|H|ψ̇⟩ = −2 Re⟨ψ|H(H−E)|ψ⟩ = −2(⟨H²⟩ − ⟨H⟩²)",
            "Monotone descent, stationary only at eigenstates",
            "Practical diagnostic: if measured E rises, suspect domain size, Δτ, ill-conditioned S, or noise",
            "Var(H) → 0 certifies an eigenstate: an eigenvalue lies within √Var of E"], y=2.7, size=20)

n += 1; s = slide("Why we cannot just run e^{−τH}", n)
bullets(s, ["A circuit implements a unitary: U†U = I",
            "(e^{−τH})† e^{−τH} = e^{−2τH} ≠ I — it rescales each eigencomponent by e^{−τEₙ}",
            "Three workarounds:",
            ("QITE: local unitary surrogate fitted to the current state (Motta et al. 2020)",),
            ("VarQITE: parametrised ansatz + McLachlan ODE for the parameters",),
            ("Probabilistic / block-encoding: ancilla + post-selection, success decays like e^{−τ}",),
            "QITE and VarQITE are the NISQ-friendly options"], size=22)

n += 1; s = slide("QITE in one slide", n)
eq(s, "e^{−Δτ H} ≈ Π_m e^{−Δτ h_m}        e^{−Δτ h_m}|ψ⟩ / √c_m  ≈  e^{−iΔτ A_m}|ψ⟩", 0.5, 1.3, 12.3, 0.9, size=24)
bullets(s, ["H = Σ_m h_m, with h_m = α_m P_m a Pauli string: e^{−Δτ h_m} = cosh(Δτα_m) − sinh(Δτα_m) P_m",
            "c_m = ⟨ψ|e^{−2Δτ h_m}|ψ⟩ is the norm that is lost",
            "A_m = Σ_μ a_μ σ_μ over a domain of Pauli strings near the support of h_m",
            "Unitary step: e^{−iΔτ A_m} = product of Pauli rotations (CNOT ladder + R_z, or native ZZ gates)",
            "Repeat over all terms m, then over steps, until E converges"], y=2.5, size=20)

n += 1; s = slide("Deriving the linear system for the coefficients a", n)
bullets(s, ["Expand to first order:  −iA|ψ⟩ ≈ |d⟩ := (|ψ′⟩ − |ψ⟩)/Δτ",
            "Minimise ‖ −iA|ψ⟩ − |d⟩ ‖² over real a"], y=1.3, h=1.6, size=20)
eq(s, "S a = b      S_μν = Re⟨ψ|σ_μ σ_ν|ψ⟩      b_μ = −Im⟨ψ|σ_μ|d⟩", 0.5, 3.0, 12.3, 0.9, size=26)
eq(s, "b_μ = sinh(Δτ α_m) / (Δτ √c_m) · Im⟨ψ|σ_μ P_m|ψ⟩", 2.0, 4.1, 9.3, 0.9, size=24)
bullets(s, ["Nonzero only when σ_μ anticommutes with P_m  ⇒  each entry is one Pauli expectation",
            "Only basis rotations + Z-measurements: no ancillas, no Hadamard test"], y=5.3, size=19)

n += 1; s = slide("The algorithm (one imaginary-time step)", n)
bullets(s, ["1.  Prepare |ψ⟩ = U_k ⋯ U₁ |ψ₀⟩   (all previous unitaries)",
            "2.  For each Hamiltonian term m:",
            ("choose domain D_m; measure ⟨h_m⟩, S_μν, b_μ on the device",),
            ("solve S a = b (regularised least squares, classical)",),
            ("append U = exp(−iΔτ Σ a_μ σ_μ) to the circuit",),
            "3.  Measure E = ⟨H⟩ and Var(H); stop when converged",
            "Quantum device: measurements + state prep.   Classical computer: small linear solves."], size=22)

n += 1; s = slide("Cost and pitfalls", n)
bullets(s, ["Circuit depth grows linearly with the number of steps: step k runs U_k ⋯ U₁|ψ₀⟩",
            ("each step is shallow, the total is not — re-preparing the prefix is part of the cost",),
            "Domain size D: strings grow like 4^D; error falls like e^{−D/ξ} only for finite correlation length ξ",
            ("molecules after fermion→qubit maps are non-local: domain can be exponential",),
            "S becomes singular near the ground state; shot noise is amplified by cond(S)",
            ("regularise / truncated SVD; adaptive Δτ",),
            "Δτ: too large → Trotter and first-order error; too small → many noisy steps"], size=21)

n += 1; s = slide("Related: variational QITE (McLachlan)", n)
eq(s, "Σ_j M_ij θ̇_j = −C_i      M_ij = Re⟨∂ᵢψ|∂ⱼψ⟩      C_i = ½ ∂ᵢE", 0.5, 1.4, 12.3, 0.9, size=26)
bullets(s, ["Equivalent to natural-gradient descent on E(θ)",
            "Fixed circuit depth (set by the ansatz) but limited by ansatz expressibility and local minima",
            "QITE has no ansatz but a growing circuit — a growing-ansatz version of the same linear system"], y=2.7, size=21)

n += 1; s = slide("QITE vs Quantum Krylov", n)
table(s, [["", "Quantum Krylov", "QITE"],
          ["Idea", "Build subspace, solve eigenproblem", "Drive state toward E₀ step by step"],
          ["Quantum work", "Overlaps ⟨ψ₀|e^{−iHt}|ψ₀⟩, ⟨ψ₀|H e^{−iHt}|ψ₀⟩", "Pauli expectations S, b at every step"],
          ["Circuit", "Fixed e^{−iHt} (Trotter), depth ∝ t", "Prefix grows with step number"],
          ["Classical work", "One small generalised eigenproblem H C = E S C", "Least-squares solve per term per step"],
          ["Main error", "Trotter; ill-conditioned S", "Domain, first-order Δτ, ill-conditioned S"],
          ["Noise", "Enters overlaps; threshold needed", "Accumulates in the state"]],
      0.5, 1.4, 12.3, [2.2, 5.2, 4.9], size=15, rowh=0.7)

n += 1; s = slide("Worked check: H₂ (STO-3G, 0.735 Å, 2 qubits)", n)
s.shapes.add_picture(str(ROOT/"figures"/"h2_convergence.png"), Inches(0.4), Inches(1.3), width=Inches(8.4))
bullets(s, ["E_FCI = −1.857275 Ha; HF error 20.3 mHa",
            "Exact S, b: chemical accuracy in 7 steps; best error 3×10⁻⁸ at τ=1.6, then settles at 1.3×10⁻⁴: Trotter error, ∝ Δτ²",
            "10³ shots/entry: error ≈0.15 Ha, worse than HF; ridge ≫ noise fixes it",
            "10⁵ shots/entry: ≈2×10⁻³ Ha",
            "Krylov exact at K=2: H₂ is a 2×2 sector, so it tests correctness, not scaling"], x=8.9, y=1.3, w=4.2, size=15)
tb(s, 0.5, 5.0, 8.2, 1.8, [("Shot noise, not algorithmic error, dominates; use ≥ 4-qubit systems (LiH, H₄) to discriminate methods.", dict(size=17, color=RED))])

n += 1; s = slide("Benchmark design on Quantinuum", n)
table(s, [["Metric", "Definition"],
          ["Accuracy", "|E − E_FCI| (mHa), Var(H) if available"],
          ["Shots / circuits", "Totals over ALL circuits incl. S, b, overlaps; distinct executions"],
          ["Depth, 2Q gates", "Max depth and total N₂q"],
          ["Qubits, runtime", "N_q (+ancilla); wall time, queue separately"]],
      0.5, 1.3, 12.3, [3.0, 9.3], size=15, rowh=0.55)
bullets(s, ["A  ideal simulation   B  shared noisy simulation   C  Quantinuum hardware",
            "D  fixed budget (shots / N₂q)  →  compare accuracy",
            "E  scaling: |E − E_FCI| vs hardware cost, vary K (Krylov) and steps N (QITE)  ← central figure"], y=4.4, size=19)

n += 1; s = slide("Fair-comparison rules and caveats", n)
bullets(s, ["Same HF start, mapping, tapering, Hamiltonian grouping",
            "Count every circuit, including repeated QITE prefixes and Hadamard-test ancillas",
            "Tune hyper-parameters (Δτ, domain, regularisation; δt, K, threshold) on a noisy simulator, then freeze",
            "Repeat over seeds; add HF and fixed-ansatz VQE baselines",
            "QITE is not better just because it is another ground-state method"], size=20, h=3.3)
tb(s, 0.6, 5.3, 12.1, 1.3, [("Research question: under identical Quantinuum resource constraints, how do QITE and Quantum Krylov trade accuracy, circuit depth, measurements and robustness?", dict(size=20, color=BLUE, bold=True))])

n += 1; s = slide("References (check before citing)", n)
bullets(s, ["Motta et al., Nature Physics 16, 205 (2020) — QITE",
            "McArdle et al., npj Quantum Information 5, 75 (2019) — variational ITE",
            "Parrish & McMahon, arXiv:1909.08925 — quantum filter diagonalisation",
            "Stair, Huang, Evangelista, JCTC 16, 2236 (2020) — multireference quantum Krylov",
            "Epperly, Lin, Nakatsukasa, SIMAX 43, 1263 (2022) — theory of quantum subspace diagonalisation",
            "Companion notes and code: quantum-imaginary-time-evolution.pdf, code/qite_h2.py"], size=19)

prs.save(ROOT / "study-notes" / "qite-slides.pptx"); print("saved", n, "slides")
